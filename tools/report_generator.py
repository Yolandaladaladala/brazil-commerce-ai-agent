from __future__ import annotations

from io import BytesIO
from typing import List, Dict, Any
import re

import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as RLImage, PageBreak, KeepTogether
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont


def _chart_png(chart: dict) -> bytes | None:
    labels = chart.get("labels") or []
    values = chart.get("values") or []
    if len(labels) < 2 or len(labels) != len(values):
        return None
    try:
        values = [float(v) for v in values]
    except Exception:
        return None

    fig, ax = plt.subplots(figsize=(7.2, 3.8))
    ctype = (chart.get("type") or "bar").lower()
    if ctype == "line":
        ax.plot(labels, values, marker="o")
    else:
        ax.bar(labels, values)
    ax.set_title(chart.get("title") or "Evidence-backed comparison")
    if chart.get("unit"):
        ax.set_ylabel(chart.get("unit"))
    ax.tick_params(axis="x", labelrotation=25)
    ax.grid(axis="y", alpha=0.2)
    fig.tight_layout()
    bio = BytesIO()
    fig.savefig(bio, format="png", dpi=160, bbox_inches="tight")
    plt.close(fig)
    bio.seek(0)
    return bio.read()


def _set_doc_fonts(doc: Document, language: str = "zh") -> None:
    styles = doc.styles
    body_font = "Aptos"
    cjk_font = "Microsoft YaHei"
    for style_name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3"]:
        style = styles[style_name]
        style.font.name = body_font
        style._element.rPr.rFonts.set(qn("w:eastAsia"), cjk_font)
    styles["Normal"].font.size = Pt(10.5)


def _add_markdown_to_docx(doc: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    for raw in lines:
        line = raw.rstrip()
        if not line:
            doc.add_paragraph()
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif re.match(r"^[-*]\s+", line):
            doc.add_paragraph(re.sub(r"^[-*]\s+", "", line), style="List Bullet")
        elif re.match(r"^\d+\.\s+", line):
            doc.add_paragraph(re.sub(r"^\d+\.\s+", "", line), style="List Number")
        elif line.startswith("|") and line.endswith("|"):
            # Keep markdown tables readable without fragile parsing.
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.font.name = "Courier New"
            run.font.size = Pt(8.5)
        else:
            text = re.sub(r"\*\*(.*?)\*\*", r"\1", line)
            doc.add_paragraph(text)


def generate_market_docx(result: Dict[str, Any], language: str = "zh") -> bytes:
    doc = Document()
    _set_doc_fonts(doc, language)

    title = f"{result.get('product','Market Research')} - {result.get('country','')}"
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(22)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Market & Product Intelligence Report").italic = True

    stats = result.get("research_stats") or {}
    doc.add_paragraph(
        f"Research scope: {result.get('objective','')}\n"
        f"Queries run: {stats.get('queries_run', 0)} | Evidence items retained: {stats.get('evidence_items', 0)}"
    )

    doc.add_page_break()
    _add_markdown_to_docx(doc, result.get("analysis") or "")

    charts = result.get("charts") or []
    for chart in charts:
        png = _chart_png(chart)
        if png:
            doc.add_heading(chart.get("title") or "Chart", level=2)
            doc.add_picture(BytesIO(png), width=Inches(6.2))
            if chart.get("note"):
                doc.add_paragraph(chart.get("note"))
            if chart.get("source_ids"):
                doc.add_paragraph("Sources: " + ", ".join(chart.get("source_ids")))

    doc.add_page_break()
    doc.add_heading("Sources & Evidence Appendix", level=1)
    for i, src in enumerate(result.get("evidence") or [], start=1):
        p = doc.add_paragraph()
        p.add_run(f"[S{i}] {src.get('title','Untitled source')}").bold = True
        if src.get("source_type"):
            p.add_run(f"\nType: {src.get('source_type')}")
        if src.get("date"):
            p.add_run(f"\nDate: {src.get('date')}")
        if src.get("link"):
            p.add_run(f"\nURL: {src.get('link')}")
        if src.get("snippet"):
            p.add_run(f"\nEvidence note: {src.get('snippet')}")

    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()


def _pdf_font(language: str) -> str:
    if language == "zh":
        try:
            pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
            return "STSong-Light"
        except Exception:
            return "Helvetica"
    return "Helvetica"


def _markdown_to_pdf_flow(markdown: str, styles, body_font: str):
    story = []
    body = ParagraphStyle("BCOSBody", parent=styles["BodyText"], fontName=body_font, fontSize=9.5, leading=14)
    h1 = ParagraphStyle("BCOSH1", parent=styles["Heading1"], fontName=body_font, fontSize=16, leading=20, spaceAfter=8)
    h2 = ParagraphStyle("BCOSH2", parent=styles["Heading2"], fontName=body_font, fontSize=13, leading=17, spaceBefore=7, spaceAfter=5)
    h3 = ParagraphStyle("BCOSH3", parent=styles["Heading3"], fontName=body_font, fontSize=11, leading=15, spaceBefore=5, spaceAfter=4)

    for raw in markdown.splitlines():
        line = raw.strip()
        if not line:
            story.append(Spacer(1, 5))
            continue
        safe = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        safe = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", safe)
        if line.startswith("### "):
            story.append(Paragraph(safe[4:], h3))
        elif line.startswith("## "):
            story.append(Paragraph(safe[3:], h2))
        elif line.startswith("# "):
            story.append(Paragraph(safe[2:], h1))
        elif re.match(r"^[-*]\s+", line):
            story.append(Paragraph("• " + re.sub(r"^[-*]\s+", "", safe), body))
        elif re.match(r"^\d+\.\s+", line):
            story.append(Paragraph(safe, body))
        else:
            story.append(Paragraph(safe, body))
    return story


def generate_market_pdf(result: Dict[str, Any], language: str = "zh") -> bytes:
    bio = BytesIO()
    body_font = _pdf_font(language)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "BCOSTitle",
        parent=styles["Title"],
        fontName=body_font,
        fontSize=20,
        leading=25,
        alignment=TA_CENTER,
        spaceAfter=10,
    )
    sub_style = ParagraphStyle(
        "BCOSSub",
        parent=styles["BodyText"],
        fontName=body_font,
        fontSize=9.5,
        leading=14,
        alignment=TA_CENTER,
        spaceAfter=8,
    )

    doc = SimpleDocTemplate(
        bio,
        pagesize=A4,
        rightMargin=42,
        leftMargin=42,
        topMargin=46,
        bottomMargin=46,
        title=f"{result.get('product','Market Research')} - {result.get('country','')}",
    )

    story = []
    story.append(Paragraph(f"{result.get('product','Market Research')} - {result.get('country','')}", title_style))
    story.append(Paragraph("Market &amp; Product Intelligence Report", sub_style))
    stats = result.get("research_stats") or {}
    story.append(Paragraph(
        f"Queries run: {stats.get('queries_run', 0)} | Evidence items retained: {stats.get('evidence_items', 0)}",
        sub_style,
    ))
    story.append(PageBreak())
    story.extend(_markdown_to_pdf_flow(result.get("analysis") or "", styles, body_font))

    for chart in result.get("charts") or []:
        png = _chart_png(chart)
        if png:
            story.append(Spacer(1, 10))
            story.append(Paragraph(chart.get("title") or "Chart", ParagraphStyle(
                "ChartTitle", parent=styles["Heading2"], fontName=body_font, fontSize=12, leading=16)))
            story.append(RLImage(BytesIO(png), width=450, height=238))
            if chart.get("note"):
                story.append(Paragraph(chart.get("note"), ParagraphStyle(
                    "ChartNote", parent=styles["BodyText"], fontName=body_font, fontSize=8.5, leading=12)))

    story.append(PageBreak())
    story.append(Paragraph("Sources &amp; Evidence Appendix", ParagraphStyle(
        "Appendix", parent=styles["Heading1"], fontName=body_font, fontSize=16, leading=20)))
    src_style = ParagraphStyle("Src", parent=styles["BodyText"], fontName=body_font, fontSize=8.2, leading=11)
    for i, src in enumerate(result.get("evidence") or [], start=1):
        title = (src.get("title") or "Untitled source").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        url = (src.get("link") or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        note = (src.get("snippet") or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        story.append(Paragraph(f"<b>[S{i}] {title}</b><br/>{url}<br/>{note}", src_style))
        story.append(Spacer(1, 4))

    doc.build(story)
    return bio.getvalue()
